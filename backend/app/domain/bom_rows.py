"""BOM intake: read CSV/XLSX into rows, then validate each row against the catalog.

Pure code. Database lookups (fuzzy suggestions) come in through CatalogIndex.suggest.
"""

import csv
import io
import re
import uuid
from collections.abc import Callable
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from fractions import Fraction
from typing import Any

from openpyxl import load_workbook

from app.domain.units import (
    Conversion,
    UnitError,
    conversion_factor,
    normalize_unit,
    to_canonical_milli,
)

MAX_ROWS = 500
MAX_FILE_BYTES = 5 * 1024 * 1024
# Bid window (1 day) + negotiation (1 day) + delivery (1 day). Earlier dates can't be met.
MIN_LEAD_DAYS = 3
WHOLE_UNITS = {"nos", "bag", "box"}

FIELDS = ("item", "spec", "quantity", "unit", "needed_by", "site", "partial_allowed", "notes")
REQUIRED = ("item", "quantity", "unit", "needed_by")
HEADER_ALIASES = {
    "item": ["item", "material", "item name", "description", "product"],
    "spec": ["spec", "grade", "grade/spec", "specification", "grade / spec"],
    "quantity": ["quantity", "qty", "qty."],
    "unit": ["unit", "uom", "units"],
    "needed_by": ["needed_by", "needed by", "required by", "date", "delivery date", "need by"],
    "site": ["site", "site name"],
    "partial_allowed": ["partial_allowed", "partial allowed", "partial", "partial ok"],
    "notes": ["notes", "note", "remarks"],
}
_HEADER_LOOKUP = {alias: key for key, aliases in HEADER_ALIASES.items() for alias in aliases}


class FileError(ValueError):
    """The whole file is unusable (empty, wrong columns, too many rows, ...)."""


def _norm(text: str) -> str:
    return " ".join(re.sub(r"[^a-z0-9./ ]+", " ", text.lower()).split())


def template_csv() -> str:
    return (
        "item,spec,quantity,unit,needed_by,partial_allowed,notes\n"
        "OPC 53 Grade Cement,OPC 53,30,bag,2026-10-15,no,\n"
        "TMT Bar Fe 500D,Fe 500D,2.5,tonne,2026-10-20,yes,8 mm and 12 mm mix\n"
    )


# --- reading ------------------------------------------------------------------------------


def _decode_csv(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        # Excel on Windows saves "CSV" as cp1252; fall back rather than reject.
        return data.decode("cp1252", errors="replace")


def _map_headers(raw: list[Any]) -> dict[int, str]:
    mapped: dict[int, str] = {}
    for i, h in enumerate(raw):
        key = _HEADER_LOOKUP.get(" ".join(str(h or "").strip().lower().split()))
        if key and key not in mapped.values():
            mapped[i] = key
    missing = [k for k in REQUIRED if k not in mapped.values()]
    if missing:
        raise FileError(
            f"Missing column(s): {', '.join(missing)}. "
            f"Expected columns: {', '.join(FIELDS)} (download the template)."
        )
    return mapped


def _rows_from_matrix(matrix: list[list[Any]]) -> list[dict[str, Any]]:
    matrix = [r for r in matrix if any(c is not None and str(c).strip() for c in r)]
    if not matrix:
        raise FileError("The file is empty.")
    headers = _map_headers(matrix[0])
    body = matrix[1:]
    if not body:
        raise FileError("The file has a header row but no items.")
    if len(body) > MAX_ROWS:
        raise FileError(f"Too many rows: {len(body)}. The limit is {MAX_ROWS} per BOM; split it.")
    return [{key: (r[i] if i < len(r) else None) for i, key in headers.items()} for r in body]


def read_table(data: bytes, filename: str) -> list[dict[str, Any]]:
    if not data:
        raise FileError("The file is empty.")
    if len(data) > MAX_FILE_BYTES:
        raise FileError("The file is larger than 5 MB.")
    name = filename.lower()
    if name.endswith(".xlsx"):
        return _read_xlsx(data)
    if name.endswith(".csv"):
        text = _decode_csv(data)
        return _rows_from_matrix([list(r) for r in csv.reader(io.StringIO(text))])
    raise FileError("Upload a .csv or .xlsx file.")


def _read_xlsx(data: bytes) -> list[dict[str, Any]]:
    try:
        values = load_workbook(io.BytesIO(data), data_only=True).worksheets[0]
        formulas = load_workbook(io.BytesIO(data), data_only=False).worksheets[0]
    except Exception as e:  # openpyxl raises many types on bad files
        raise FileError("This is not a readable .xlsx file.") from e
    matrix: list[list[Any]] = []
    for vrow, frow in zip(values.iter_rows(), formulas.iter_rows(), strict=True):
        cells = []
        for v, f in zip(vrow, frow, strict=True):
            if v.value is None and isinstance(f.value, str) and f.value.startswith("="):
                # A formula with no saved result (file never opened in Excel).
                cells.append(_Uncomputed(f.value))
            else:
                cells.append(v.value)
        matrix.append(cells)
    return _rows_from_matrix(matrix)


@dataclass(frozen=True)
class _Uncomputed:
    formula: str

    def __str__(self) -> str:
        return self.formula


# --- validation ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ItemInfo:
    id: uuid.UUID
    code: str
    name: str
    grade: str | None
    canonical_unit: str
    aliases: tuple[str, ...]


@dataclass
class CatalogIndex:
    items: list[ItemInfo]
    conversions: dict[uuid.UUID | None, list[Conversion]]
    suggest: Callable[[str], list[ItemInfo]]
    by_term: dict[str, ItemInfo] = field(init=False)

    def __post_init__(self) -> None:
        self.by_term = {}
        for it in self.items:
            for term in (it.name, it.code, *it.aliases):
                self.by_term.setdefault(_norm(term), it)

    def convs(self, item: ItemInfo) -> list[Conversion]:
        return self.conversions.get(item.id, []) + self.conversions.get(None, [])

    def get(self, item_id: uuid.UUID) -> ItemInfo | None:
        return next((i for i in self.items if i.id == item_id), None)


@dataclass
class RowResult:
    line_no: int
    input: dict[str, str]
    item: ItemInfo | None = None
    suggestions: list[ItemInfo] = field(default_factory=list)
    qty_canonical_milli: int | None = None
    needed_by: date | None = None
    partial_allowed: bool = False
    errors: list[dict[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    merged_into: int | None = None

    def err(self, fld: str, msg: str) -> None:
        self.errors.append({"field": fld, "message": msg})

    def as_dict(self) -> dict[str, Any]:
        return {
            "line_no": self.line_no,
            "input": self.input,
            "item": _item_out(self.item) if self.item else None,
            "suggestions": [_item_out(s) for s in self.suggestions],
            "qty_canonical_milli": self.qty_canonical_milli,
            "canonical_unit": self.item.canonical_unit if self.item else None,
            "needed_by": self.needed_by.isoformat() if self.needed_by else None,
            "partial_allowed": self.partial_allowed,
            "errors": self.errors,
            "warnings": self.warnings,
            "merged_into": self.merged_into,
        }


def _item_out(i: ItemInfo) -> dict[str, Any]:
    return {"id": str(i.id), "code": i.code, "name": i.name, "grade": i.grade,
            "canonical_unit": i.canonical_unit}  # fmt: skip


def cell_text(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


_DATE_FORMATS = ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%d.%m.%Y", "%d %b %Y", "%d %B %Y", "%d-%b-%Y")


def parse_date(text: str) -> date | None:
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(text.strip(), fmt).date()
        except ValueError:
            continue
    return None


def parse_qty(text: str) -> Fraction | None:
    try:
        d = Decimal(text.replace(",", "").strip())
    except InvalidOperation:
        return None
    return Fraction(d) if d.is_finite() else None


_YES = {"yes", "y", "true", "1"}
_NO = {"no", "n", "false", "0", ""}


def _match_item(index: CatalogIndex, item_text: str, spec_text: str) -> ItemInfo | None:
    for candidate in (f"{item_text} {spec_text}", item_text, f"{spec_text} {item_text}"):
        hit = index.by_term.get(_norm(candidate))
        if hit:
            return hit
    return None


def validate_row(
    line_no: int,
    raw: dict[str, Any],
    index: CatalogIndex,
    *,
    today: date,
    site_name: str,
    other_site_names: set[str],
) -> RowResult:
    inp = {k: cell_text(raw.get(k)) for k in FIELDS}
    r = RowResult(line_no=line_no, input=inp)

    for k, v in raw.items():
        if isinstance(v, _Uncomputed):
            r.err(
                k,
                f"Formula {v.formula} has no saved value. Open the file in Excel, save, re-upload.",
            )
    if r.errors:
        return r

    # item
    chosen = raw.get("catalog_item_id")
    if chosen:
        try:
            r.item = index.get(uuid.UUID(str(chosen)))
        except ValueError:
            r.item = None
        if r.item is None:
            r.err("item", "Chosen catalog item does not exist.")
    elif not inp["item"]:
        r.err("item", "Item is required.")
    else:
        r.item = _match_item(index, inp["item"], inp["spec"])
        if r.item is None:
            r.suggestions = index.suggest(f"{inp['item']} {inp['spec']}".strip())[:3]
            hint = (
                f" Did you mean: {', '.join(s.name for s in r.suggestions)}?"
                if r.suggestions
                else ""
            )
            r.err("item", f"Unknown item '{inp['item']}'. Choose a catalog item.{hint}")

    # quantity
    qty = parse_qty(inp["quantity"]) if inp["quantity"] else None
    if not inp["quantity"]:
        r.err("quantity", "Quantity is required.")
    elif qty is None:
        r.err("quantity", f"Quantity '{inp['quantity']}' is not a number.")
    elif qty <= 0:
        r.err("quantity", "Quantity must be more than zero.")

    # unit
    unit = None
    if not inp["unit"]:
        r.err("unit", "Unit is required.")
    else:
        try:
            unit = normalize_unit(inp["unit"])
        except UnitError:
            r.err("unit", f"Unit '{inp['unit']}' is not recognised." + _unit_hint(r.item, index))

    if r.item and unit and qty is not None and qty > 0:
        try:
            conversion_factor(unit, r.item.canonical_unit, index.convs(r.item))
        except UnitError:
            r.err(
                "unit",
                f"'{inp['unit']}' cannot be converted for {r.item.name}."
                + _unit_hint(r.item, index),
            )
        else:
            try:
                milli = to_canonical_milli(qty, unit, r.item.canonical_unit, index.convs(r.item))
            except UnitError:
                r.err("quantity", "Quantity has too many decimal places.")
            else:
                if r.item.canonical_unit in WHOLE_UNITS and milli % 1000:
                    r.err("quantity", f"Must be a whole number of {r.item.canonical_unit}.")
                else:
                    r.qty_canonical_milli = milli

    # needed-by
    if not inp["needed_by"]:
        r.err("needed_by", "Needed-by date is required.")
    else:
        d = parse_date(inp["needed_by"])
        earliest = today + timedelta(days=MIN_LEAD_DAYS)
        if d is None:
            r.err(
                "needed_by",
                f"Date '{inp['needed_by']}' not understood. Use YYYY-MM-DD or DD/MM/YYYY.",
            )
        elif d < today:
            r.err("needed_by", "Needed-by date is in the past.")
        elif d < earliest:
            r.err("needed_by", f"Too soon. Earliest feasible date is {earliest:%d %b %Y} "
                               "(bid window, negotiation and delivery).")  # fmt: skip
        else:
            r.needed_by = d

    # site
    if inp["site"]:
        s = inp["site"].strip().lower()
        if s in other_site_names:
            r.err(
                "site", f"This BOM is for {site_name}. Put {inp['site']} lines in a separate BOM."
            )
        elif s != site_name.lower():
            r.err("site", f"Unknown site '{inp['site']}'.")

    # partial
    p = inp["partial_allowed"].strip().lower()
    if p in _YES:
        r.partial_allowed = True
    elif p not in _NO:
        r.err("partial_allowed", "Use yes or no.")
    return r


def _unit_hint(item: ItemInfo | None, index: CatalogIndex) -> str:
    if item is None:
        return ""
    units = {item.canonical_unit}
    for c in index.convs(item):
        for u in (c.from_unit, c.to_unit):
            try:
                conversion_factor(u, item.canonical_unit, index.convs(item))
                units.add(u)
            except UnitError:
                pass
    return f" Use one of: {', '.join(sorted(units))}."


def merge_duplicates(rows: list[RowResult]) -> None:
    """Same item, needed-by date and partial flag -> quantities added onto the first line."""
    first: dict[tuple[Any, ...], RowResult] = {}
    for r in rows:
        if r.errors or r.item is None or r.qty_canonical_milli is None:
            continue
        key = (r.item.id, r.needed_by, r.partial_allowed)
        if key in first:
            keep = first[key]
            assert keep.qty_canonical_milli is not None
            keep.qty_canonical_milli += r.qty_canonical_milli
            keep.warnings.append(f"Line {r.line_no} merged into this line (same item and date).")
            r.merged_into = keep.line_no
            r.warnings.append(f"Duplicate of line {keep.line_no}; merged into it.")
        else:
            first[key] = r


def validate_rows(
    raws: list[dict[str, Any]],
    index: CatalogIndex,
    *,
    today: date,
    site_name: str,
    other_site_names: set[str],
) -> list[RowResult]:
    if not raws:
        raise FileError("Add at least one line.")
    if len(raws) > MAX_ROWS:
        raise FileError(f"Too many rows: {len(raws)}. The limit is {MAX_ROWS} per BOM; split it.")
    rows = [
        validate_row(i + 1, raw, index, today=today, site_name=site_name,
                     other_site_names=other_site_names)
        for i, raw in enumerate(raws)
    ]  # fmt: skip
    merge_duplicates(rows)
    return rows
